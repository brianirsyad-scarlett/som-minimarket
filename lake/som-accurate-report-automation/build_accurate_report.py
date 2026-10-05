"""
Build "0. 2025 Accurate.xlsx" (table Accurate_Report) - what Sell In reads - from the
monthly Accurate exports, without Excel.

Locally this is two Power Query refreshes:

  Accurate 2026.xlsx     [table Accurate_2026]  <- every "*Sales*.xlsx" in the folder
                                                   (sheet "Delivery Order Detail"),
                                                   minus returns ("*Return*.xlsx",
                                                   sheet "Sales Return Detail"), cleaned
  0. 2025 Accurate.xlsx  [table Accurate_Report] <- Accurate_2025 + Accurate_2026

Both M queries are repeated here step by step (see the comments). Accurate 2025 is a
frozen table and is taken as-is from "Accurate 2025.xlsx".

    python build_accurate_report.py --source-dir <folder of monthly exports> \
        --accurate-2025 "Accurate 2025.xlsx" --out "0. 2025 Accurate.xlsx"
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
import warnings
from pathlib import Path

import fastexcel
import polars as pl
import python_calamine as pc
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

SALES_SHEET = "Delivery Order Detail"
RETURN_SHEET = "Sales Return Detail"

SALES_COLUMNS = ["CreatedOn", "SentOn", "SO Number", "DO Number", "Sales Type",
                 "Customer", "City", "Item Name", "Quantity", "Price"]
RETURN_COLUMNS = ["Date", "DO Number", "Return Number", "Customer", "Item Name", "Quantity"]

# Accurate_2026's column order (M: Accurate Report, after the join and Quantity column).
YEAR_COLUMNS = ["CreatedOn", "SentOn", "SO Number", "DO Number", "Sales Type", "Customer",
                "City", "Item Name", "Qty Sales", "Revenue", "Qty Return", "Quantity"]
# Accurate_Report's column order (Accurate_2025 first, so Province sits after Customer).
REPORT_COLUMNS = ["CreatedOn", "SentOn", "SO Number", "DO Number", "Sales Type", "Customer",
                  "Province", "City", "Item Name", "Qty Sales", "Revenue", "Qty Return", "Quantity"]
DATE_COLUMNS = {"CreatedOn", "SentOn"}
INT_COLUMNS = {"Qty Sales", "Revenue", "Qty Return", "Quantity"}

# M: #"Cleaned Sales Type" - plain text replacements, applied in this order.
SALES_TYPE_REPLACEMENTS = [
    ("01. ", ""), ("02. ", ""), ("03. ", ""), ("04. ", ""), ("05. ", ""), ("06. ", ""),
    ("07. ", ""), ("08. ", ""), ("09. ", ""), ("10. ", ""), ("11. ", ""),
    ("Ecommerce - Tiktok", "TikTok Shop"),
    ("Ecommerce - ", ""),
]


# --------------------------------------------------------------------------- #
# Reading the exports - Excel.Workbook([Content], true) + ExpandTableColumn
# --------------------------------------------------------------------------- #

def _cell(v):
    # Power Query reads an empty cell as null.
    return None if v == "" else v


def read_sheet(path: Path, sheet: str, columns: list[str]) -> list[dict]:
    """The named columns of one sheet, first row promoted to headers. A file
    without the sheet contributes nothing, as the M's Name.1 filter does."""
    wb = pc.CalamineWorkbook.from_path(str(path))
    if sheet not in wb.sheet_names:
        return []
    rows = wb.get_sheet_by_name(sheet).to_python(skip_empty_area=False)
    if not rows:
        return []
    header = [str(h) if h is not None else "" for h in rows[0]]
    idx = {c: header.index(c) for c in columns if c in header}
    out = []
    for r in rows[1:]:
        out.append({c: (_cell(r[idx[c]]) if c in idx and idx[c] < len(r) else None) for c in columns})
    return out


def _to_date(v):
    """type date. Anything that is not a date becomes null (an error cell in M,
    which these rows never survive: they are the repeated header rows)."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d %b %Y"):
            try:
                return dt.datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _to_int(v):
    """Int64.Type: rounds to the nearest whole number."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    try:
        return int(round(float(str(v).replace(",", ""))))
    except ValueError:
        return None


def _to_text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def source_files(source_dir: Path, word: str) -> list[Path]:
    """M: Folder.Files (recursive) filtered on Text.Contains([Name], word) - case-sensitive."""
    return sorted((p for p in source_dir.rglob("*.xlsx")
                   if word in p.name and not p.name.startswith("~$")),
                  key=lambda p: (str(p.parent).lower(), p.name.lower()))


# --------------------------------------------------------------------------- #
# M: Accurate Sales / Accurate Return / Accurate Report
# --------------------------------------------------------------------------- #

def accurate_sales(source_dir: Path) -> pl.DataFrame:
    recs = []
    for f in source_files(source_dir, "Sales"):
        for r in read_sheet(f, SALES_SHEET, SALES_COLUMNS):
            recs.append({
                "CreatedOn": _to_date(r["CreatedOn"]), "SentOn": _to_date(r["SentOn"]),
                "SO Number": _to_text(r["SO Number"]), "DO Number": _to_text(r["DO Number"]),
                "Sales Type": _to_text(r["Sales Type"]), "Customer": _to_text(r["Customer"]),
                "City": _to_text(r["City"]), "Item Name": _to_text(r["Item Name"]),
                "Qty Sales": _to_int(r["Quantity"]), "Revenue": _to_int(r["Price"]),
            })
    schema = {"CreatedOn": pl.Date, "SentOn": pl.Date, "SO Number": pl.String,
              "DO Number": pl.String, "Sales Type": pl.String, "Customer": pl.String,
              "City": pl.String, "Item Name": pl.String, "Qty Sales": pl.Int64, "Revenue": pl.Int64}
    return pl.DataFrame(recs, schema=schema, orient="row") if recs else pl.DataFrame(schema=schema)


def accurate_return(source_dir: Path) -> pl.DataFrame:
    recs = []
    for f in source_files(source_dir, "Return"):
        for r in read_sheet(f, RETURN_SHEET, RETURN_COLUMNS):
            if r["Date"] == "Date":      # M: #"Filtered Rows2" - repeated header rows
                continue
            recs.append({"DO Number": _to_text(r["DO Number"]), "Customer": _to_text(r["Customer"]),
                         "Item Name": _to_text(r["Item Name"]), "Qty Return": _to_int(r["Quantity"])})
    schema = {"DO Number": pl.String, "Customer": pl.String, "Item Name": pl.String, "Qty Return": pl.Int64}
    return pl.DataFrame(recs, schema=schema, orient="row") if recs else pl.DataFrame(schema=schema)


def accurate_year(source_dir: Path) -> pl.DataFrame:
    """M: Accurate Report - the Accurate_2026 table."""
    sales, returns = accurate_sales(source_dir), accurate_return(source_dir)
    df = sales.join(returns, on=["DO Number", "Customer", "Item Name"], how="left",
                    nulls_equal=True, maintain_order="left")
    df = df.with_columns(pl.col("Qty Return").fill_null(0))
    df = df.with_columns((pl.col("Qty Sales") - pl.col("Qty Return")).alias("Quantity"))

    st = pl.col("Sales Type")
    for old, new in SALES_TYPE_REPLACEMENTS:
        st = st.str.replace_all(old, new, literal=True)
    df = df.with_columns(st.alias("Sales Type"))
    df = df.with_columns([pl.col(c).str.to_uppercase() for c in ("Sales Type", "Customer", "City", "Item Name")])

    df = df.filter(pl.col("Sales Type").is_not_null()
                   & (pl.col("Customer").ne("PT OPTO LUMBUNG SEJAHTERA") | pl.col("Customer").is_null()))
    # Table.Distinct on [DO Number]&[Item Name]: first row per key; a null in either
    # part makes the whole key null, and all null keys count as one.
    key = pl.when(pl.col("DO Number").is_null() | pl.col("Item Name").is_null()).then(None) \
            .otherwise(pl.col("DO Number") + pl.col("Item Name"))
    df = df.with_columns(key.alias("_key")).unique(subset="_key", keep="first", maintain_order=True).drop("_key")
    df = df.filter(~pl.col("Customer").str.contains("OPTO LINGKAR", literal=True).fill_null(False))
    df = df.filter(pl.col("Sales Type") != "SALES TYPE")
    return df.select(YEAR_COLUMNS)


def accurate_2025(path: Path) -> pl.DataFrame:
    df = fastexcel.read_excel(str(path)).load_table("Accurate_2025").to_polars()
    return normalise(df)


def normalise(df: pl.DataFrame) -> pl.DataFrame:
    """The Changed Type steps: dates as Date, counts as Int64, the rest text."""
    exprs = []
    for c in df.columns:
        t = df.schema[c]
        if c in DATE_COLUMNS:
            if isinstance(t, pl.Datetime):
                exprs.append(pl.col(c).cast(pl.Date))
            elif t == pl.String:
                exprs.append(pl.col(c).str.to_date("%d/%m/%Y", strict=False))
        elif c in INT_COLUMNS:
            if t != pl.Int64:
                exprs.append(pl.col(c).cast(pl.Float64, strict=False).round(0).cast(pl.Int64))
        elif t != pl.String:
            exprs.append(pl.col(c).cast(pl.String))
    return df.with_columns(exprs) if exprs else df


def accurate_report(y2025: pl.DataFrame, years: list[pl.DataFrame]) -> pl.DataFrame:
    """M: Table.Combine({Accurate_2025, Accurate_2026, ...})."""
    return pl.concat([y2025, *years], how="diagonal_relaxed").select(REPORT_COLUMNS)


# --------------------------------------------------------------------------- #
# Writing - a real Excel table, since Sell In reads it with load_table("Accurate_Report")
# --------------------------------------------------------------------------- #

def write_table_xlsx(df: pl.DataFrame, path: Path, table_name: str, sheet: str) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet)
    cols = df.columns
    ws.append(cols)
    date_idx = [i for i, c in enumerate(cols) if c in DATE_COLUMNS]
    for row in df.iter_rows():
        cells = list(row)
        for i in date_idx:
            if cells[i] is not None:
                cell = WriteOnlyCell(ws, value=cells[i])
                cell.number_format = "dd/mm/yyyy"
                cells[i] = cell
        ws.append(cells)
    if df.height == 0:
        ws.append([None] * len(cols))
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(cols))}{max(df.height, 1) + 1}")
    table.tableColumns = [TableColumn(id=i, name=n) for i, n in enumerate(cols, start=1)]
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=".*table columns manually.*")
        ws.add_table(table)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source-dir", type=Path, action="append", required=True,
                    help="Folder of monthly exports for one year; repeat for several years.")
    ap.add_argument("--accurate-2025", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args(argv)

    years = []
    for d in args.source_dir:
        y = accurate_year(d)
        print(f"{d}: {y.height:,} rows")
        years.append(y)
    y2025 = accurate_2025(args.accurate_2025)
    print(f"Accurate_2025: {y2025.height:,} rows")
    report = accurate_report(y2025, years)
    write_table_xlsx(report, args.out, "Accurate_Report", "Accurate_Report")
    print(f"wrote {args.out}: {report.height:,} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
