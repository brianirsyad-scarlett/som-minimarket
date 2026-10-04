"""
National Distribution quarter workbooks, built in the cloud (no Excel needed).

The laptop's "National Distribution Raw_<from>_<to>.xlsx" are Power Query workbooks: they read the
compiled Distributor table, rename / upper-case a few columns, add Region and Store Type Group, join
the Product master and keep one quarter of dates. This module does the same steps in pandas and
writes a plain workbook whose first sheet is that table - exactly what the Summary workbook's own
query reads (it takes the first sheet of every "National Distribution Raw*" file).

    input    gs://bucket_som/sales_parquet/Distributor_Sales.parquet         (built by sell_through.py)
             gs://bucket_som/sales_parquet/raw/master data/Master Data Sales.xlsx   (table "Product")
    output   gs://bucket_som/sales_parquet/raw/sell_through/National Distribution Raw_<from>_<to>.xlsx

    python national_distribution.py run                   # gate -> build -> validate -> back up -> publish
    python national_distribution.py run --force
    python national_distribution.py build --parquet P --master M --outdir DIR   # local, no GCS
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd

WIB = ZoneInfo("Asia/Jakarta")
BUCKET = "bucket_som"
PRODUCTION_PARQUET = "sales_parquet/Distributor_Sales.parquet"
MASTER = "sales_parquet/raw/master data/Master Data Sales.xlsx"
OUT_PREFIX = "sales_parquet/raw/sell_through/"
DRAFT_PREFIX = "sales_parquet/raw/sell_through/cloud/"
BACKUP_PREFIX = "sales_parquet/raw/sell_through/cloud/backup/"
KEEP_BACKUPS = 3                      # per workbook; the files are ~12 MB

def quarters(today: dt.date | None = None) -> list[tuple[dt.date, dt.date]]:
    """(first day, last day) of the previous, the current and the next calendar quarter, as of today (WIB).
    A workbook that does not exist yet is created; older quarters are history and stay as they are."""
    today = today or dt.datetime.now(WIB).date()
    base = today.year * 4 + (today.month - 1) // 3
    out = []
    for i in (base - 1, base, base + 1):
        year, q = divmod(i, 4)
        first = dt.date(year, q * 3 + 1, 1)
        following = dt.date(year, q * 3 + 4, 1) if q < 3 else dt.date(year + 1, 1, 1)
        out.append((first, following - dt.timedelta(days=1)))
    return out


SHEET = "National Sell Through by Area"
SOURCE_COLS = ["RSP/Distributor Name", "Date", "CustomerName", "Store Type", "ItemName", "SE", "ASS",
               "City", "Total", "SourceFile", "SheetName"]          # first 11 columns, as the query's Columns=11
OUT_COLS = ["Area", "RSP/Distributor Name", "Date", "CustomerName", "Store Type", "ItemName", "Salesman",
            "ASS", "City", "Quantity", "Region", "Store Type Group",
            "Brand", "Category", "Sub Category", "Variant", "Product Name", "Type of Item"]
UPPER_COLS = ["Area", "RSP/Distributor Name", "CustomerName", "Store Type", "ItemName", "Salesman", "ASS", "City"]
MASTER_COLS = ["Brand", "Category", "Sub Category", "Variant", "Product Name", "Type of Item"]

MIN_ROWS_VS_PREVIOUS = 0.90           # vs the row count recorded on our last build of the same file
MIN_BRAND_MATCH = 0.80                # share of rows that found their product in the master


def quarter_name(a: dt.date, b: dt.date) -> str:
    return f"National Distribution Raw_{a:%Y%m%d}_{b:%Y%m%d}.xlsx"


# --------------------------------------------------------------------------- #
# The Power Query steps
# --------------------------------------------------------------------------- #

def read_master_raw(path: Path) -> pd.DataFrame:
    """The "Product" table exactly as in the workbook (No dropped, ItemName trimmed + upper-cased)."""
    with zipfile.ZipFile(path) as z:
        ref = None
        for n in z.namelist():
            if re.match(r"xl/tables/table\d+\.xml$", n):
                x = z.read(n).decode("utf-8", "ignore")
                if re.search(r'displayName="Product"', x):
                    ref = re.search(r' ref="([^"]+)"', x).group(1)
        if ref is None:
            raise SystemExit('Master Data Sales.xlsx has no table named "Product"')
    first, last = (int(re.search(r"\d+", p).group()) for p in ref.split(":"))
    c1, c2 = re.findall(r"[A-Z]+", ref.replace("$", ""))
    m = pd.read_excel(path, sheet_name="Product", header=first - 1, nrows=last - first,
                      usecols=f"{c1}:{c2}", engine="openpyxl")
    m.columns = [str(c).strip() for c in m.columns]
    m = m.drop(columns=["No"])
    m["ItemName"] = m["ItemName"].astype("string").str.strip().str.upper()
    return m


def load_master(path: Path) -> pd.DataFrame:
    """'Master Product' query: first row per ItemName."""
    m = read_master_raw(path)
    return m.drop_duplicates(subset=["ItemName"], keep="first")[["ItemName"] + MASTER_COLS].reset_index(drop=True)


def quantity(series: pd.Series) -> pd.Series:
    """Power Query's Int64.Type on the Total column (round half to even, like Number.Round's default)."""
    return pd.to_numeric(series, errors="coerce").round(0).astype("Int64")


def transform(dist: pd.DataFrame) -> pd.DataFrame:
    """'National_Distribution (Control)' query."""
    df = dist[SOURCE_COLS].copy()
    df["Date"] = pd.to_datetime(df["Date"], format="%Y-%m-%d", errors="coerce")
    df["Total"] = quantity(df["Total"])
    df["SourceFile"] = df["SourceFile"].astype("string").str.replace(" Sell Through.xlsx", "", regex=False)
    df = df.rename(columns={"SE": "Salesman", "Total": "Quantity", "SourceFile": "Area"})
    df = df.drop(columns=["SheetName"])
    for c in UPPER_COLS:
        # Power Query reads a blank CSV field as an empty string, so Excel holds "" (not a blank cell) here
        df[c] = df[c].astype("string").str.upper().fillna("")
    area = df["Area"].fillna("")
    df["Region"] = "WEST"
    df.loc[area.str.contains("6.", regex=False) | area.str.contains("7.", regex=False)
           | area.str.contains("8.", regex=False), "Region"] = "EAST"
    df["Store Type Group"] = "GT"
    df.loc[df["Store Type"].fillna("").str.contains("MT", regex=False), "Store Type Group"] = "MT"
    return df


def quarter_table(df: pd.DataFrame, master: pd.DataFrame, first: dt.date, last: dt.date) -> pd.DataFrame:
    """'National Sell Through by Area' query: join the master, keep the quarter's dates."""
    out = df.merge(master, how="left", on="ItemName", sort=False)
    keep = (out["Date"] >= pd.Timestamp(first)) & (out["Date"] <= pd.Timestamp(last))
    return out.loc[keep, OUT_COLS].reset_index(drop=True)


# --------------------------------------------------------------------------- #
# Workbook
# --------------------------------------------------------------------------- #

def write_workbook(table: pd.DataFrame, path: Path) -> None:
    """One sheet + an Excel table. Written cell by cell because pandas turns "" into a blank cell, and
    Excel's own file holds real empty strings in the text columns (blank only where the master had no match)."""
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path), {"strings_to_numbers": False, "strings_to_formulas": False,
                                         "strings_to_urls": False})
    ws = wb.add_worksheet(SHEET)
    date_fmt = wb.add_format({"num_format": "yyyy-mm-dd"})
    ws.write_row(0, 0, OUT_COLS)
    for c, col in enumerate(OUT_COLS):
        values = table[col]
        if col == "Date":
            for r, v in enumerate(values, start=1):
                if pd.notna(v):
                    ws.write_datetime(r, c, v.to_pydatetime(), date_fmt)
        elif col == "Quantity":
            for r, v in enumerate(values, start=1):
                if pd.notna(v):
                    ws.write_number(r, c, int(v))
        elif col in MASTER_COLS:
            for r, v in enumerate(values.tolist(), start=1):
                if v is not None and not pd.isna(v):
                    ws.write_string(r, c, v)
        else:
            for r, v in enumerate(values.tolist(), start=1):
                ws.write_string(r, c, "" if v is None or pd.isna(v) else v)
    if len(table):
        ws.add_table(0, 0, len(table), len(OUT_COLS) - 1, {
            "name": "National_Sell_Through_by_Area", "style": "TableStyleMedium2",
            "columns": [{"header": c} for c in OUT_COLS]})
    ws.set_column(0, len(OUT_COLS) - 1, 18)
    wb.close()


def build_all(parquet: Path, master_path: Path, outdir: Path, qs: list) -> dict[str, pd.DataFrame]:
    dist = pd.read_parquet(parquet)
    master = load_master(master_path)
    print(f"distributor rows {len(dist):,} | master items {len(master):,}")
    df = transform(dist)
    built = {}
    for first, last in qs:
        tbl = quarter_table(df, master, first, last)
        name = quarter_name(first, last)
        write_workbook(tbl, outdir / name)
        built[name] = tbl
        print(f"  {name}: {len(tbl):,} rows, qty {tbl['Quantity'].sum():,}")
    return built


# --------------------------------------------------------------------------- #
# Checks: do the PUBLISHED workbooks equal the parquet?  which master items do not match?
# --------------------------------------------------------------------------- #

def _normalise(t: pd.DataFrame) -> pd.DataFrame:
    t = t.dropna(how="all").copy()                      # Excel keeps one blank row for an empty table
    if not len(t):
        return t
    t["Date"] = pd.to_datetime(t["Date"])
    t["Quantity"] = pd.to_numeric(t["Quantity"]).astype("int64")
    return t.sort_values(OUT_COLS, na_position="first", kind="stable").reset_index(drop=True)


def mismatched_cells(a: pd.DataFrame, b: pd.DataFrame) -> int:
    a, b = _normalise(a), _normalise(b)
    if len(a) != len(b):
        return -1
    return int(sum(int((~((a[c] == b[c]) | (a[c].isna() & b[c].isna()))).sum()) for c in OUT_COLS))


def reconcile(dist: pd.DataFrame, qs: list, built: dict, fetch) -> list[str]:
    """Raw parquet -> rows and quantity inside each quarter, compared with the workbook that was
    actually uploaded (fetch(name) returns it as a DataFrame) and with the table we built."""
    d = pd.to_datetime(dist["Date"], format="%Y-%m-%d", errors="coerce")
    qty = pd.to_numeric(dist["Total"], errors="coerce").round(0)
    problems, in_any = [], pd.Series(False, index=dist.index)
    print(f"{'workbook':50s} {'parquet rows':>13s} {'xlsx rows':>10s} {'qty equal':>10s} {'cell diffs':>10s}")
    for first, last in qs:
        name = quarter_name(first, last)
        inq = (d >= pd.Timestamp(first)) & (d <= pd.Timestamp(last))
        in_any |= inq
        pub = _normalise(fetch(name))
        got_rows = len(pub)
        got_qty = int(pub["Quantity"].sum()) if got_rows else 0
        diffs = mismatched_cells(built[name], pub)
        print(f"{name:50s} {int(inq.sum()):>13,} {got_rows:>10,} {str(got_qty == int(qty[inq].sum())):>10s} {diffs:>10}")
        if got_rows != int(inq.sum()) or got_qty != int(qty[inq].sum()) or diffs != 0:
            problems.append(f"{name}: published workbook differs from the parquet "
                            f"(rows {got_rows:,} vs {int(inq.sum()):,}, qty {got_qty:,} vs {int(qty[inq].sum()):,}, cell diffs {diffs})")
    print(f"parquet rows outside these quarters: {int((d.notna() & ~in_any).sum()):,} (older quarters) | "
          f"rows with no Date: {int(d.isna().sum()):,} (never in any workbook)")
    return problems


def missing_items(built: dict, master_raw: pd.DataFrame) -> pd.DataFrame:
    """Items that were sold but have no row in the master (their Brand..Type of Item stay blank)."""
    norm = lambda x: re.sub(r"[^A-Z0-9]", "", str(x).upper())
    keys = {norm(k): k for k in master_raw["ItemName"].dropna()}
    sales = pd.concat(built.values(), ignore_index=True)
    un = sales[sales["Brand"].isna()]
    if not len(un):
        return pd.DataFrame(columns=["ItemName", "rows", "qty", "areas", "last_date", "looks_like_master_item"])
    g = (un.groupby("ItemName").agg(rows=("Quantity", "size"), qty=("Quantity", "sum"),
         areas=("Area", lambda s: ", ".join(sorted(set(s.dropna()))[:3])), last_date=("Date", "max"))
         .sort_values("qty", ascending=False).reset_index())
    g["last_date"] = g["last_date"].dt.strftime("%Y-%m-%d")
    g["looks_like_master_item"] = g["ItemName"].map(lambda x: keys.get(norm(x), ""))
    return g


MISSING_STATE = "_state/som-sellthrough/master_missing.sha256"


def format_missing_email(missing: pd.DataFrame, run_url: str) -> tuple[str, str]:
    n = len(missing)
    subject = f"[Sell Through] {n} item{'s' if n != 1 else ''} sold but missing in Master Data Sales"
    near = int((missing["looks_like_master_item"] != "").sum())
    lines = [f"{n} item(s) in the Sell Through workbooks have no row in the Product table of Master Data Sales.xlsx,",
             "so Brand / Category / Sub Category / Variant / Product Name / Type of Item are blank for them",
             f"({int(missing['qty'].sum()):,} qty, {int(missing['rows'].sum()):,} rows in the previous / current / next quarter).",
             "",
             f"{near} of them already exist in the master with a slightly different spelling (spacing, dashes): fix the name",
             "in the regional workbook or add the exact sales spelling to the master. The rest are new items to add.",
             "", "ItemName (as sold)  |  rows  |  qty  |  areas  |  last date  |  looks like master item", "-" * 100]
    for r in missing.itertuples(index=False):
        lines.append(f"{r.ItemName}  |  {r.rows}  |  {r.qty:,}  |  {r.areas}  |  {r.last_date}  |  {r.looks_like_master_item or '-'}")
    lines += ["", "Add the missing rows to the Product table in Master Data Sales.xlsx (sales_parquet/raw/master data/).",
              "The next run picks them up automatically. This email is sent once per change of this list.", run_url]
    return subject, "\n".join(lines)


def email_missing(bucket, missing: pd.DataFrame, run_url: str, send=None) -> None:
    """Email the missing items once per change of the list (not every 30 minutes). Nothing is written to the bucket
    except a hash of the last list sent."""
    import hashlib
    if send is None:
        from notify import send
    digest = hashlib.sha256("\n".join(sorted(missing["ItemName"])).encode()).hexdigest() if len(missing) else "none"
    state = bucket.blob(MISSING_STATE)
    last = state.download_as_text().strip() if state.exists() else ""
    if digest == last:
        print(f"master check: {len(missing)} item(s) missing - same list as the last email, not sent again")
        return
    if not len(missing):
        state.upload_from_string(digest)
        print("master check: nothing missing")
        return
    subject, body = format_missing_email(missing, run_url)
    with tempfile.TemporaryDirectory() as tmp:
        attach = Path(tmp) / "items_missing_in_master.csv"
        missing.to_csv(attach, index=False)
        try:
            send(subject, body, attach)
        except SystemExit as e:                 # no SMTP secrets yet: stay silent in the log, retry next time
            print(f"master check: {len(missing)} item(s) missing but the email was NOT sent ({e})")
            return
    state.upload_from_string(digest)
    print(f"master check: {len(missing)} item(s) missing - emailed")


# --------------------------------------------------------------------------- #
# GCS
# --------------------------------------------------------------------------- #

def _bucket():
    from google.cloud import storage
    return storage.Client().bucket(BUCKET)


def gate(bucket, force: bool, qs: list) -> tuple[bool, str]:
    parquet, master = bucket.get_blob(PRODUCTION_PARQUET), bucket.get_blob(MASTER)
    if parquet is None or master is None:
        return False, f"missing input: parquet={parquet is not None} master={master is not None}"
    outs = [bucket.get_blob(OUT_PREFIX + quarter_name(a, b)) for a, b in qs]
    newest_input = max(parquet.updated, master.updated)
    print(f"inputs: parquet {parquet.updated.astimezone(WIB):%Y-%m-%d %H:%M}, master {master.updated.astimezone(WIB):%Y-%m-%d %H:%M} WIB")
    if force:
        return True, "forced"
    if any(o is None for o in outs):
        return True, "a quarter workbook is missing"
    oldest_output = min(o.updated for o in outs)
    print(f"oldest quarter workbook: {oldest_output.astimezone(WIB):%Y-%m-%d %H:%M} WIB")
    if newest_input > oldest_output:
        return True, "an input is newer than the workbooks"
    return False, "workbooks are newer than both inputs"


def validate(name: str, tbl: pd.DataFrame, first: dt.date, last: dt.date, previous) -> list[str]:
    problems = []
    if list(tbl.columns) != OUT_COLS:
        problems.append(f"{name}: unexpected columns {list(tbl.columns)}")
    if len(tbl):
        lo, hi = tbl["Date"].min().date(), tbl["Date"].max().date()
        if lo < first or hi > last:
            problems.append(f"{name}: dates {lo}..{hi} outside {first}..{last}")
        brand = tbl["Brand"].notna().mean()
        if brand < MIN_BRAND_MATCH:
            problems.append(f"{name}: only {brand:.1%} of rows matched the product master")
    prev_rows = int((previous.metadata or {}).get("rows", 0)) if previous is not None else 0
    if prev_rows and len(tbl) < MIN_ROWS_VS_PREVIOUS * prev_rows:
        problems.append(f"{name}: {len(tbl):,} rows is more than 10% below the last cloud build's {prev_rows:,}")
    return problems


def publish(bucket, local: Path, name: str, rows: int, qty: int, parquet_updated) -> None:
    final, draft = OUT_PREFIX + name, DRAFT_PREFIX + name
    blob = bucket.blob(draft)
    blob.metadata = {"rows": str(rows), "qty": str(qty), "built_from_parquet": parquet_updated.isoformat()}
    blob.upload_from_filename(str(local), timeout=900,
                              content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    old = bucket.get_blob(final)
    if old is not None:
        stamp = dt.datetime.now(WIB).strftime("%Y%m%d-%H%M")
        stem = name.rsplit(".", 1)[0]
        bucket.copy_blob(old, bucket, f"{BACKUP_PREFIX}{stem}-{stamp}.xlsx")
        for stale in sorted(bucket.list_blobs(prefix=BACKUP_PREFIX + stem + "-"), key=lambda b: b.name)[:-KEEP_BACKUPS]:
            stale.delete()
    bucket.copy_blob(blob, bucket, final)
    print(f"published {final} ({rows:,} rows)")


def run(args) -> int:
    bucket = _bucket()
    qs = quarters()
    print("quarters:", ", ".join(f"{a:%Y-%m-%d}..{b:%Y-%m-%d}" for a, b in qs))
    go, why = gate(bucket, args.force, qs)
    print(f"build: {go} ({why})")
    if not go:
        return 0
    parquet = bucket.get_blob(PRODUCTION_PARQUET)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        parquet.download_to_filename(str(tmp / "dist.parquet"), timeout=900)
        bucket.blob(MASTER).download_to_filename(str(tmp / "master.xlsx"), timeout=900)
        built = build_all(tmp / "dist.parquet", tmp / "master.xlsx", tmp, qs)
        problems = []
        for first, last in qs:
            name = quarter_name(first, last)
            problems += validate(name, built[name], first, last, bucket.get_blob(OUT_PREFIX + name))
        if problems:
            print("NOT published:")
            for p in problems:
                print("  -", p)
            return 1
        if args.no_publish:
            return 0
        for first, last in qs:
            name = quarter_name(first, last)
            tbl = built[name]
            publish(bucket, tmp / name, name, len(tbl), int(tbl["Quantity"].sum()), parquet.updated)

        # Items sold but missing in the master -> email (never printed: this repo's logs are public)
        email_missing(bucket, missing_items(built, read_master_raw(tmp / "master.xlsx")),
                      os.environ.get("RUN_URL", ""))

        # Re-read what is now in the bucket and prove it equals the parquet
        def fetch(name: str) -> pd.DataFrame:
            from openpyxl import load_workbook
            local = tmp / ("published_" + name)
            bucket.blob(OUT_PREFIX + name).download_to_filename(str(local), timeout=900)
            wb = load_workbook(local, read_only=True, data_only=True)
            rows = wb[SHEET].iter_rows(values_only=True)
            header = list(next(rows))
            df = pd.DataFrame(list(rows), columns=header)
            wb.close()
            return df

        problems = reconcile(pd.read_parquet(tmp / "dist.parquet"), qs, built, fetch)
        if problems:
            print("RECONCILIATION FAILED:")
            for p in problems:
                print("  -", p)
            return 1
        print("reconciliation OK: every published workbook equals the parquet")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--force", action="store_true")
    r.add_argument("--no-publish", action="store_true")
    b = sub.add_parser("build")
    b.add_argument("--parquet", type=Path, required=True)
    b.add_argument("--master", type=Path, required=True)
    b.add_argument("--outdir", type=Path, required=True)
    a = ap.parse_args(argv)
    if a.cmd == "build":
        a.outdir.mkdir(parents=True, exist_ok=True)
        build_all(a.parquet, a.master, a.outdir, quarters())
        return 0
    return run(a)


if __name__ == "__main__":
    sys.exit(main())
