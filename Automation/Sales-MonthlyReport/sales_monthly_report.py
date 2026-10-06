#!/usr/bin/env python
"""
Rebuilds the monthly sales workbook that Power Query used to produce by hand,
e.g. Data\\Report\\Sales\\2026\\2026 09 Sep.xlsx

It is a line-by-line port of the M queries embedded in that workbook
(customXml/item1.xml -> Formulas/Section1.m):

    E_Stock_Report      <- E-Stock 2025.xlsx            [table E_Stock_2025]
    Odoo_Report         <- Odoo Report.xlsx             [table Odoo_Report]
    Accurate_Report     <- 0. 2025 Accurate.xlsx        [table Accurate_Report]
    Anchanto_Report(Qn) <- Anchanto Report <Y> Q<n>.xlsx [every table/sheet]
    Master Product      <- Master Data Sales.xlsx       [sheet Product, skip 3]
    Master Region       <- Master Data Sales.xlsx       [sheet Area,    skip 5]
    Master Channel      <- Master Data Sales.xlsx       [sheet Channel]

    <month> Offline     = Combine(E-Stock, Odoo, Accurate)
                          + Master Product / Region / Channel lookups
                          filtered to the month, Channel Group <> "ONLINE"
    <month> Online (n)  = same base + Anchanto appended,
                          filtered to a day-range of the month, Channel Group = "ONLINE"

The output has to stay a real Excel workbook with real ListObjects (tables),
because the yearly roll-up (Data\\Report\\Sales\\<year>\\<year> Report.xlsx)
reads this folder with Folder.Files + Excel.Workbook and keeps only Kind="Table".
A sheet without a table is invisible to it.

Usage
-----
    python sales_monthly_report.py                     # current month
    python sales_monthly_report.py --month 2026-09
    python sales_monthly_report.py --month prev
    python sales_monthly_report.py --month 2026-07 --month 2026-08   # backfill
    python sales_monthly_report.py --month 2026-09 --dry-run         # counts only
"""

from __future__ import annotations

import argparse
import calendar
import datetime as dt
import logging
import os
import re
import shutil
import sys
import time
import warnings
from pathlib import Path
from typing import Iterable, Sequence

import fastexcel
import polars as pl
import python_calamine as pc
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

# --------------------------------------------------------------------------- #
# Paths
# --------------------------------------------------------------------------- #

HERE = Path(__file__).resolve().parent
SOM_ROOT = HERE.parents[1]                       # ...\SCARLETT-329\SOM
DATA = SOM_ROOT / "Data"

SALES_DIR = DATA / "Report" / "Sales"
P_ESTOCK = SALES_DIR / "E-Stock 2025" / "E-Stock 2025.xlsx"
P_ODOO = SALES_DIR / "Odoo" / "Odoo Report.xlsx"
P_ACCURATE = SALES_DIR / "Accurate Report 2025" / "0. 2025 Accurate.xlsx"
P_ANCHANTO_DIR = SALES_DIR / "Anchanto Report"
P_MASTER = DATA / "Master Data Sales" / "Matrix" / "Master Data Sales.xlsx"

# Backups and the staging file deliberately live OUTSIDE Data\Report\Sales\<year>.
# The yearly roll-up globs that folder and only skips names containing
# "$", "Report" or ".ini" - a stray "*.bak.xlsx" there would be counted twice.
TMP_DIR = HERE / ".tmp"
BACKUP_DIR = HERE / "backups"
# Newest backups kept per month workbook (~100 MB each). Unbounded, this grew by
# ~3 GB a month - September alone left 31 copies - before 2026-10-06.
BACKUP_KEEP = 3
LOG_FILE = HERE / "run.log"

MONTH_ABBR = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

EXCEL_MAX_ROWS = 1_048_576
DATE_FMT = "dd/mm/yyyy"

# Final column order + the exact header spelling the yearly roll-up expects.
OUTPUT_COLUMNS: list[str] = [
    "CreatedOn", "SentOn", "NoPesanan", "SalesType", "SubSalesType",
    "Channel Group", "Channel", "Sub Channel", "Region", "Province", "City",
    "CustomerName", "Payment Terms", "Customer & Address", "InvoiceOn",
    "SO Number", "Qty Sales", "Revenue", "Qty Return", "ReportSource",
    "ItemName", "Quantity", "Status", "Brand", "Category", "SubCategory",
    "Variant", "Product Name", "Type of Item",
]
DATE_COLUMNS = {"CreatedOn", "SentOn", "InvoiceOn"}

log = logging.getLogger("sales_monthly_report")


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #

def _join(left: pl.DataFrame, right: pl.DataFrame, on: Sequence[str]) -> pl.DataFrame:
    """Left join that matches null==null, the way Power Query's Table.NestedJoin does."""
    try:
        return left.join(right, on=list(on), how="left", nulls_equal=True)
    except TypeError:  # polars < 1.24
        return left.join(right, on=list(on), how="left", join_nulls=True)


def _as_date(df: pl.DataFrame, *names: str) -> pl.DataFrame:
    """Power Query's `type date`: drop any time component, parse text timestamps."""
    exprs = []
    for name in names:
        if name not in df.columns:
            continue
        dtype = df.schema[name]
        if dtype == pl.Date:
            continue
        if isinstance(dtype, pl.Datetime):
            exprs.append(pl.col(name).cast(pl.Date).alias(name))
        elif dtype == pl.String:
            # Explicit formats: polars raises rather than returning nulls when it
            # cannot infer one, and these columns are mixed date/blank/number.
            exprs.append(
                pl.coalesce(
                    pl.col(name).str.to_datetime(
                        format="%Y-%m-%d %H:%M:%S", strict=False).cast(pl.Date),
                    pl.col(name).str.to_date(format="%Y-%m-%d", strict=False),
                    pl.col(name).str.to_date(format="%d/%m/%Y", strict=False),
                ).alias(name)
            )
        # numeric columns are left alone - M does not coerce those either
    return df.with_columns(exprs) if exprs else df


def _as_int(df: pl.DataFrame, *names: str) -> pl.DataFrame:
    """Power Query's Int64.Type (rounds, unlike a plain truncating cast)."""
    exprs = [
        pl.col(n).round(0).cast(pl.Int64, strict=False).alias(n)
        for n in names
        if n in df.columns and df.schema[n] != pl.Int64
    ]
    return df.with_columns(exprs) if exprs else df


def _upper(df: pl.DataFrame, *names: str) -> pl.DataFrame:
    """Text.Upper - null stays null."""
    exprs = [
        pl.col(n).cast(pl.String).str.to_uppercase().alias(n)
        for n in names
        if n in df.columns
    ]
    return df.with_columns(exprs) if exprs else df


def _text(df: pl.DataFrame, *names: str) -> pl.DataFrame:
    exprs = [
        pl.col(n).cast(pl.String).alias(n)
        for n in names
        if n in df.columns and df.schema[n] != pl.String
    ]
    return df.with_columns(exprs) if exprs else df


def _load_table(path: Path, table: str) -> pl.DataFrame:
    t0 = time.perf_counter()
    df = fastexcel.read_excel(str(path)).load_table(table).to_polars()
    log.info("  read table %-18s %8d rows  (%.1fs)  %s",
             table, df.height, time.perf_counter() - t0, path.name)
    return df


def _cell_text(v) -> str | None:
    """How Power Query stringifies a raw sheet cell when it types the column as text."""
    if v is None:
        return None
    if isinstance(v, str):
        return v or None
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return str(v)


def _sheet_rows(path: Path, sheet: str) -> list[list]:
    wb = pc.CalamineWorkbook.from_path(str(path))
    return wb.get_sheet_by_name(sheet).to_python(skip_empty_area=False)


def _text_frame(rows: Iterable[Sequence], colmap: dict[str, int]) -> pl.DataFrame:
    data: dict[str, list] = {name: [] for name in colmap}
    for row in rows:
        for name, idx in colmap.items():
            data[name].append(_cell_text(row[idx]) if idx < len(row) else None)
    return pl.DataFrame(data, schema={name: pl.String for name in colmap})


# --------------------------------------------------------------------------- #
# Master tables
# --------------------------------------------------------------------------- #

def master_product() -> pl.DataFrame:
    """M: sheet 'Product', Table.Skip 3, Distinct on ItemName, then Text.Upper."""
    rows = _sheet_rows(P_MASTER, "Product")[3:]
    df = _text_frame(rows, {
        "ItemName": 1, "Brand": 2, "Category": 3, "SubCategory": 4,
        "Variant": 5, "Product Name": 6, "Type of Item": 7,
    })
    # Distinct happens BEFORE the uppercase in the M, so names differing only by
    # case survive it and can become duplicate keys afterwards. Kept faithful,
    # but warn - duplicate keys multiply rows in the lookup.
    df = df.unique(subset=["ItemName"], keep="first", maintain_order=True)
    df = _upper(df, "ItemName")
    dupes = df.height - df.select("ItemName").unique().height
    if dupes:
        log.warning("  Master Product has %d duplicate ItemName keys after "
                    "uppercasing - lookups will duplicate those rows", dupes)
    log.info("  read Master Product     %8d rows", df.height)
    return df


def master_region() -> pl.DataFrame:
    """M: sheet 'Area', Table.Skip 5, Distinct on Province, drop Province = "0"."""
    rows = _sheet_rows(P_MASTER, "Area")[5:]
    df = _text_frame(rows, {"City": 1, "Region": 2, "Province": 3})
    df = df.unique(subset=["Province"], keep="first", maintain_order=True)
    df = df.filter(pl.col("Province").ne("0") | pl.col("Province").is_null())
    df = df.select("Province", "Region")
    log.info("  read Master Region      %8d rows", df.height)
    return df


def master_channel() -> pl.DataFrame:
    """M: sheet 'Channel', Table.PromoteHeaders."""
    rows = _sheet_rows(P_MASTER, "Channel")
    header = [str(h) for h in rows[0]]
    idx = {name: header.index(name) for name in
           ("Channel Group", "SalesType", "SubSalesType", "Channel", "Sub Channel")}
    df = _text_frame(rows[1:], idx)
    df = df.select("SalesType", "SubSalesType", "Channel Group", "Channel", "Sub Channel")
    dupes = df.height - df.select("SalesType", "SubSalesType").unique().height
    if dupes:
        log.warning("  Master Channel has %d duplicate (SalesType, SubSalesType) "
                    "keys - lookups will duplicate those rows", dupes)
    log.info("  read Master Channel     %8d rows", df.height)
    return df


# --------------------------------------------------------------------------- #
# Source queries
# --------------------------------------------------------------------------- #

def q_e_stock() -> pl.DataFrame:
    df = _load_table(P_ESTOCK, "E_Stock_2025")
    df = df.with_columns(pl.lit("E-Stock").alias("ReportSource"))
    df = _as_date(df, "CreatedOn", "SentOn")
    df = _text(df, "City", "NoPesanan", "CustomerName")
    df = _as_int(df, "Quantity")
    df = _upper(df, "SalesType", "City", "NoPesanan", "CustomerName",
                "ItemName", "ReportSource")
    return df.select("CreatedOn", "SentOn", "SalesType", "City", "NoPesanan",
                     "CustomerName", "ItemName", "Quantity", "ReportSource")


def q_odoo() -> pl.DataFrame:
    df = _load_table(P_ODOO, "Odoo_Report")
    df = df.with_columns(pl.lit("Odoo_OptoLumbung").alias("ReportSource"))
    df = _as_date(df, "CreatedOn", "SentOn")
    df = _as_int(df, "Quantity")
    df = _as_date(df, "InvoiceOn")          # not typed in M, but it is a date column
    df = df.with_columns(
        pl.col("CustomerState").str.replace_all(" (ID)", "", literal=True)
    )

    # M: `not Text.Contains([Checked], "Reported")`.
    # Text.Contains is case-sensitive, and on a null [Checked] the whole
    # expression is null, which Table.SelectRows drops. Ported exactly.
    keep = ~pl.col("Checked").str.contains("Reported", literal=True)
    df = df.filter(keep.fill_null(False))

    df = df.drop([c for c in (
        "CustomerEntity", "SalesPerson", "Qty Ordered", "Qty To Deliver",
        "Qty To Invoice", "Qty Invoiced", "Status", "Invoice Status",
        "Qty Delivered", "DO Number", "SO Number", "Address", "Checked",
    ) if c in df.columns])
    df = df.rename({
        "CustomerState": "Province",
        "CustomerCity": "City",
        "Delivery Status": "Status",
        "OrderReference": "NoPesanan",
        "CustomerLevel1": "SalesType",
        "CustomerLevel2": "SubSalesType",
    })
    df = _upper(df, "SalesType", "SubSalesType", "Province", "City",
                "CustomerName", "ItemName", "ReportSource", "NoPesanan", "Status")
    df = df.filter(pl.col("SentOn") >= dt.date(2025, 1, 1))
    return df.select("CreatedOn", "SalesType", "SubSalesType", "Province", "City",
                     "CustomerName", "ItemName", "NoPesanan", "SentOn", "Status",
                     "Payment Terms", "Quantity", "Customer & Address", "InvoiceOn",
                     "ReportSource")


def q_accurate() -> pl.DataFrame:
    df = _load_table(P_ACCURATE, "Accurate_Report")
    df = df.rename({
        "DO Number": "NoPesanan",
        "Item Name": "ItemName",
        "Customer": "CustomerName",
        "Sales Type": "SalesType",
    })
    df = _upper(df, "ItemName", "City", "CustomerName", "SalesType", "NoPesanan")
    df = df.with_columns(pl.lit("ACCURATE").alias("ReportSource"))
    df = _as_date(df, "CreatedOn", "SentOn")
    df = _as_int(df, "Quantity")
    df = _text(df, "Province")
    # M: `not Text.Contains([CustomerName], "EMPLOYE")` - null CustomerName drops out.
    keep = ~pl.col("CustomerName").str.contains("EMPLOYE", literal=True)
    df = df.filter(keep.fill_null(False))
    return df.select("CreatedOn", "SentOn", "SO Number", "NoPesanan", "SalesType",
                     "CustomerName", "Province", "City", "ItemName", "Qty Sales",
                     "Revenue", "Qty Return", "Quantity", "ReportSource")


ANCHANTO_COLUMNS = [
    "CreatedOn", "SentOn", "SalesType", "NoPesanan", "ItemName", "Quantity",
    "Unit Price", "Province", "City", "Channel Group", "Channel", "Sub Channel",
    "Brand", "Category", "Sub Category", "Variant", "Product Name", "Type of Item",
]


def q_anchanto(paths: Sequence[Path], start: dt.date, end: dt.date) -> pl.DataFrame:
    """Every table (or, if the file has none, every sheet) of the given Anchanto files.

    The M expands each workbook's tables and only then filters on SentOn. We filter
    per sheet instead, purely so a 300 MB quarter file never has to sit in memory
    whole; the result is identical.
    """
    parts: list[pl.DataFrame] = []
    for path in paths:
        if not path.exists():
            log.warning("  Anchanto file not found, skipped: %s", path)
            continue
        reader = fastexcel.read_excel(str(path))
        tables = list(reader.table_names())
        if tables:
            loaders = [(name, lambda n=name: reader.load_table(n)) for name in tables]
        else:
            # The M keeps only Kind="Table". A workbook that was rewritten by a
            # script loses its ListObjects and would silently contribute nothing,
            # so fall back to its sheets and say so loudly.
            log.warning("  %s has NO Excel tables - falling back to its %d sheets. "
                        "Power Query would have read 0 rows from this file.",
                        path.name, len(reader.sheet_names))
            loaders = [(name, lambda n=name: reader.load_sheet(n, header_row=0))
                       for name in reader.sheet_names]

        for name, load in loaders:
            t0 = time.perf_counter()
            try:
                df = load().to_polars()
            except BaseException as exc:  # pyo3 PanicException derives from BaseException
                # calamine panics on a header-only table (A1:R1) - e.g. the empty
                # Nov/Dec tables of a freshly started quarter. That is zero rows.
                if "invalid range bounds" not in str(exc):
                    raise
                log.info("    %-28s %8d rows kept  (header-only table)  %s",
                         name, 0, path.name)
                continue
            df = _as_date(df, "CreatedOn", "SentOn")
            df = df.filter(pl.col("SentOn").is_between(start, end))
            if df.height:
                df = _as_int(df, "Quantity")
                df = df.drop([c for c in ("Unit Price",) if c in df.columns])
                if "Sub Category" in df.columns:
                    df = df.rename({"Sub Category": "SubCategory"})
                df = df.with_columns(
                    pl.col("SalesType").cast(pl.String).alias("CustomerName"),
                    pl.lit("ANCHANTO").alias("ReportSource"),
                )
                df = _upper(df, "ReportSource")
                parts.append(df)
            log.info("    %-28s %8d rows kept  (%.1fs)  %s",
                     name, df.height, time.perf_counter() - t0, path.name)

    if not parts:
        return pl.DataFrame(schema={c: pl.String for c in OUTPUT_COLUMNS})
    return pl.concat(parts, how="diagonal_relaxed")


# --------------------------------------------------------------------------- #
# The combined base table
# --------------------------------------------------------------------------- #

def build_base() -> pl.DataFrame:
    """Table.Combine({E_Stock, Odoo, Accurate}) + the three master lookups."""
    log.info("Reading offline sources")
    base = pl.concat([q_e_stock(), q_odoo(), q_accurate()], how="diagonal_relaxed")
    log.info("  combined                %8d rows", base.height)

    log.info("Reading master data")
    base = _join(base, master_product(), ["ItemName"])
    base = _join(base, master_region(), ["Province"])
    base = _join(base, master_channel(), ["SalesType", "SubSalesType"])
    log.info("  after lookups           %8d rows", base.height)
    return base


def shape_output(df: pl.DataFrame) -> pl.DataFrame:
    """Reorder to OUTPUT_COLUMNS, adding anything the source never supplied."""
    missing = [c for c in OUTPUT_COLUMNS if c not in df.columns]
    if missing:
        df = df.with_columns([pl.lit(None, dtype=pl.String).alias(c) for c in missing])
    return df.select(OUTPUT_COLUMNS)


# --------------------------------------------------------------------------- #
# Month plumbing
# --------------------------------------------------------------------------- #

class MonthSpec:
    def __init__(self, year: int, month: int, split_days: Sequence[int],
                 out_dir: Path | None = None):
        self.year = year
        self.month = month
        self.out_dir = out_dir
        self.last_day = calendar.monthrange(year, month)[1]
        self.start = dt.date(year, month, 1)
        self.end = dt.date(year, month, self.last_day)
        self.split_days = [d for d in split_days if 0 < d < self.last_day]

    @property
    def label(self) -> str:
        return f"{self.year} {self.month:02d} {MONTH_ABBR[self.month]}"

    @property
    def filename(self) -> str:
        return f"{self.label}.xlsx"

    @property
    def out_path(self) -> Path:
        return (self.out_dir or SALES_DIR / str(self.year)) / self.filename

    def online_ranges(self) -> list[tuple[dt.date, dt.date]]:
        bounds = [0, *self.split_days, self.last_day]
        return [
            (dt.date(self.year, self.month, lo + 1), dt.date(self.year, self.month, hi))
            for lo, hi in zip(bounds, bounds[1:])
        ]

    def anchanto_files(self) -> list[Path]:
        """The quarter this month sits in, plus the one before it - what the M does."""
        quarter = (self.month - 1) // 3 + 1
        wanted = [(self.year, quarter)]
        wanted.append((self.year, quarter - 1) if quarter > 1 else (self.year - 1, 4))
        out = []
        for year, q in wanted:
            for name in (f"Anchanto Report {year} Q{q}.xlsx", f"Anchanto Report {year}.xlsx"):
                p = P_ANCHANTO_DIR / name
                if p.exists() and p not in out:
                    out.append(p)
        return out


def parse_month(token: str, today: dt.date) -> tuple[int, int]:
    token = token.strip().lower()
    if token in ("current", "this", ""):
        return today.year, today.month
    if token in ("prev", "previous", "last"):
        first = today.replace(day=1)
        prev = first - dt.timedelta(days=1)
        return prev.year, prev.month
    m = re.fullmatch(r"(\d{4})[-_ /]?(\d{1,2})", token)
    if not m:
        raise argparse.ArgumentTypeError(
            f"cannot read month {token!r} - use YYYY-MM, 'current' or 'prev'")
    year, month = int(m.group(1)), int(m.group(2))
    if not 1 <= month <= 12:
        raise argparse.ArgumentTypeError(f"month out of range in {token!r}")
    return year, month


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #

def table_name_for(sheet_name: str) -> str:
    """'2026 09 Sep Online (1)' -> '_2026_09_Sep_Online__1', matching Excel."""
    name = re.sub(r"\W", "_", sheet_name).rstrip("_")
    return name if re.match(r"^[A-Za-z_]", name) else "_" + name


def write_workbook(path: Path, sheets: Sequence[tuple[str, pl.DataFrame]]) -> None:
    wb = Workbook(write_only=True)
    for sheet_name, df in sheets:
        if df.height + 1 > EXCEL_MAX_ROWS:
            raise SystemExit(
                f"'{sheet_name}' would need {df.height:,} data rows, over Excel's "
                f"{EXCEL_MAX_ROWS:,} row limit. Add another --split-day so the "
                f"Online data is cut into more sheets."
            )
        ws = wb.create_sheet(sheet_name)
        ws.append(OUTPUT_COLUMNS)

        date_idx = [i for i, c in enumerate(OUTPUT_COLUMNS) if c in DATE_COLUMNS]
        t0 = time.perf_counter()
        for row in df.iter_rows():
            cells = list(row)
            for i in date_idx:
                if cells[i] is not None:
                    cell = WriteOnlyCell(ws, value=cells[i])
                    cell.number_format = DATE_FMT
                    cells[i] = cell
            ws.append(cells)

        if df.height == 0:
            # Excel will not accept a header-only table. Power Query left one
            # blank row in this case (the original "Online (2)" was A1:AC2), and
            # the yearly roll-up already tolerates an all-null row.
            ws.append([None] * len(OUTPUT_COLUMNS))

        last_col = get_column_letter(len(OUTPUT_COLUMNS))
        table = Table(
            displayName=table_name_for(sheet_name),
            ref=f"A1:{last_col}{max(df.height, 1) + 1}",
        )
        table.tableColumns = [
            TableColumn(id=i, name=name) for i, name in enumerate(OUTPUT_COLUMNS, start=1)
        ]
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        with warnings.catch_warnings():
            # openpyxl always warns "you must add table columns manually" in
            # write-only mode; tableColumns is set just above, so it is noise.
            warnings.filterwarnings("ignore", message=".*table columns manually.*")
            ws.add_table(table)
        log.info("  sheet %-26s %8d rows  (%.1fs)",
                 sheet_name, df.height, time.perf_counter() - t0)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def publish(staged: Path, target: Path, backup: bool) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    if backup and target.exists():
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        kept = BACKUP_DIR / f"{target.stem} {stamp}{target.suffix}"
        shutil.copy2(target, kept)
        log.info("Backed up previous file -> %s", kept)
        # The stamp sorts chronologically, so name order is age order. Matching
        # on the exact "<stem> YYYYMMDD-HHMMSS<suffix>" shape leaves every other
        # file in backups/ (the yearly .bak files, anchanto/) alone.
        shape = re.compile(re.escape(target.stem) + r" \d{8}-\d{6}" + re.escape(target.suffix) + "$")
        mine = sorted((p for p in BACKUP_DIR.iterdir() if p.is_file() and shape.match(p.name)),
                      key=lambda p: p.name, reverse=True)
        for stale in mine[BACKUP_KEEP:]:
            stale.unlink()
            log.info("Pruned old backup %s", stale.name)
    try:
        try:
            # Atomic when staging and target share a volume, which they do for the
            # real output path (both on the SOM drive).
            os.replace(staged, target)
        except OSError as exc:
            if getattr(exc, "winerror", None) != 17:   # 17 = cross-volume move
                raise
            if target.exists():
                target.unlink()
            shutil.move(str(staged), str(target))
    except PermissionError as exc:
        raise SystemExit(
            f"Cannot replace {target} - it is probably open in Excel. "
            f"Close it and run again. The new file is waiting at {staged}"
        ) from exc


# --------------------------------------------------------------------------- #
# One month
# --------------------------------------------------------------------------- #

def build_month(spec: MonthSpec, base: pl.DataFrame, dry_run: bool, backup: bool) -> None:
    log.info("=" * 72)
    log.info("Building %s  (%s .. %s)", spec.label, spec.start, spec.end)

    in_month = pl.col("SentOn").is_between(spec.start, spec.end)
    is_online = pl.col("Channel Group") == "ONLINE"

    offline = shape_output(base.filter(in_month & ~is_online.fill_null(False)))
    log.info("  %-28s %8d rows", "Offline", offline.height)

    anchanto = q_anchanto(spec.anchanto_files(), spec.start, spec.end)
    log.info("  %-28s %8d rows", "Anchanto (in month)", anchanto.height)

    # Table.Combine first, filter second - and via `diagonal_relaxed`, so a column
    # one side lacks is filled with nulls of the other side's dtype rather than
    # dragging numeric columns up to String.
    frames = [base] + ([anchanto] if anchanto.height else [])
    online_all = shape_output(
        pl.concat(frames, how="diagonal_relaxed").filter(in_month & is_online)
    )
    log.info("  %-28s %8d rows", "Online (all)", online_all.height)

    sheets: list[tuple[str, pl.DataFrame]] = [(f"{spec.label} Offline", offline)]
    ranges = spec.online_ranges()
    for n, (lo, hi) in enumerate(ranges, start=1):
        part = online_all.filter(pl.col("SentOn").is_between(lo, hi))
        name = f"{spec.label} Online ({n})" if len(ranges) > 1 else f"{spec.label} Online"
        log.info("  %-28s %8d rows  (%s .. %s)", name, part.height, lo, hi)
        sheets.append((name, part))

    if dry_run:
        log.info("Dry run - nothing written.")
        return

    TMP_DIR.mkdir(parents=True, exist_ok=True)
    staged = TMP_DIR / spec.filename
    t0 = time.perf_counter()
    write_workbook(staged, sheets)
    log.info("Wrote %s (%.1f MB) in %.0fs",
             staged.name, staged.stat().st_size / 1e6, time.perf_counter() - t0)
    publish(staged, spec.out_path, backup)
    log.info("Published -> %s", spec.out_path)


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #

def setup_logging(verbose: bool) -> None:
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s %(message)s", "%Y-%m-%d %H:%M:%S")
    root = logging.getLogger()
    root.setLevel(logging.DEBUG if verbose else logging.INFO)
    stream = logging.StreamHandler(sys.stdout)
    stream.setFormatter(fmt)
    root.addHandler(stream)
    file_handler = logging.FileHandler(LOG_FILE, encoding="utf-8")
    file_handler.setFormatter(fmt)
    root.addHandler(file_handler)


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--month", action="append", default=[],
                    help="YYYY-MM, 'current' or 'prev'. Repeat to build several. "
                         "Default: current month.")
    ap.add_argument("--split-day", action="append", type=int, default=[],
                    help="Day of month to cut the Online sheets at. Default: 15, "
                         "giving 'Online (1)' = 1..15 and 'Online (2)' = 16..end.")
    ap.add_argument("--out-dir", type=Path, default=None,
                    help="Write somewhere other than Data\\Report\\Sales\\<year>. "
                         "Useful for a trial run - the yearly roll-up globs the real "
                         "folder, so never leave spare copies of the workbook there.")
    ap.add_argument("--no-backup", action="store_true",
                    help="Do not copy the previous workbook into ./backups first.")
    ap.add_argument("--dry-run", action="store_true",
                    help="Report the row counts but write nothing.")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)

    setup_logging(args.verbose)
    today = dt.date.today()
    months = [parse_month(m, today) for m in (args.month or ["current"])]
    split_days = args.split_day or [15]

    for path in (P_ESTOCK, P_ODOO, P_ACCURATE, P_MASTER):
        if not path.exists():
            log.error("Source missing: %s", path)
            return 2

    started = time.perf_counter()
    log.info("#" * 72)
    log.info("sales_monthly_report - months=%s split_days=%s",
             ", ".join(f"{y}-{m:02d}" for y, m in months), split_days)

    base = build_base()
    for year, month in months:
        build_month(MonthSpec(year, month, split_days, args.out_dir), base,
                    args.dry_run, not args.no_backup)

    log.info("Done in %.0fs", time.perf_counter() - started)
    return 0


if __name__ == "__main__":
    sys.exit(main())
